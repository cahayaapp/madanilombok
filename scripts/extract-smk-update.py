"""Read the user-designated second sheet; preserve merges and canonical Madani IDs."""
import json,re,hashlib
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
SOURCE=Path('/Users/haimac/Downloads/DATABASE SMKS ISLAM PLUS AL-MADANI 2026-2027.xlsx')
d=json.loads((ROOT/'seed/master-data.json').read_text());w=openpyxl.load_workbook(SOURCE,data_only=True);s=w.worksheets[1]
year='TA-2026-2027-GANJIL';update_id='smk-2026-2027-sheet2-v1'
classes={6:'CLS-SMK-X-BUSANA',7:'CLS-SMK-X-DKV',8:'CLS-SMK-XI-BUSANA',9:'CLS-SMK-XI-DKV',10:'CLS-SMK-XI-TJKT',11:'CLS-SMK-XII-DKV',12:'CLS-SMK-XII-BUSANA'}
teachers={}
for sid,staff in d['staff'].items():
 m=re.search(r'Kode Guru SMK (\d+)',staff.get('roleSummary',''))
 if m:
  assert int(m[1]) not in teachers
  teachers[int(m[1])]=sid
merges={}
for area in s.merged_cells.ranges:
 for row in range(area.min_row,area.max_row+1):
  for col in range(area.min_col,area.max_col+1):merges[row,col]=area
changes={};day=None;warnings=[];count=0;source_cells=[]
for row in range(11,108):
 if s.cell(row,3).value in ['SENIN','SELASA','RABU','KAMIS','JUMAT','SABTU']:day=s.cell(row,3).value.title()
 time=str(s.cell(row,5).value or '');m=re.fullmatch(r'(\d{2})\.(\d{2})\s*[-.]\s*(\d{2})\.(\d{2})',time)
 if not m:continue
 start=f'{m[1]}:{m[2]}';end=f'{m[3]}:{m[4]}'
 if time=='14.20.15.00':warnings.append('E89 menggunakan titik pemisah; dibaca 14.20–15.00 sesuai dua waktu yang tertulis.')
 for col,cid in classes.items():
  area=merges.get((row,col));anchor=s.cell(area.min_row,area.min_col) if area else s.cell(row,col)
  code=re.fullmatch(r'([A-U])\.(\d+)',str(anchor.value or '').strip())
  if not code:continue
  subject=f'MPL-SMK-{code[1]}';staff=teachers[int(code[2])];assert subject in d['subjects'] and cid in d['classes']
  combined=[classes[c] for c in range(area.min_col,area.max_col+1) if c in classes] if area else [cid]
  sid=f'JSMK-2627G-{row:03d}-{col:02d}'
  rec={'academicYearId':year,'unitId':'UNIT-SMK','classId':cid,'subjectId':subject,'teacherStaffId':staff,'teacherAssignmentStatus':'confirmed','day':day,'startTime':start,'endTime':end,'status':'active','source':SOURCE.name,'sourceSheet':s.title,'sourceCell':s.cell(row,col).coordinate,'sourceAnchor':anchor.coordinate,'sourceCode':code[0],'validationStatus':'Sesuai sheet kedua','seeded':True}
  if len(combined)>1:rec.update(combinedClassIds=combined,teachingSessionId=f'JSMK-2627G-GABUNG-{row:03d}-{area.min_col:02d}')
  changes[f'schedules/academic/{sid}']=rec;count+=1
# Keep prior schedule records, if any, instead of deleting attendance history.
for sid,rec in d['schedules']['academic'].items():
 if rec.get('unitId')=='UNIT-SMK' and rec.get('academicYearId')==year and not sid.startswith('JSMK-2627G-'):
  changes[f'schedules/academic/{sid}/status']='inactive';changes[f'schedules/academic/{sid}/supersededBy']=update_id
for staff in set(r['teacherStaffId'] for r in changes.values() if isinstance(r,dict)):
 taught=[r for r in changes.values() if isinstance(r,dict) and r.get('teacherStaffId')==staff]
 roles=list(d['staff'][staff].get('appRoles',[]))
 if any(r['subjectId']!='MPL-SMK-U' for r in taught):roles.append('guru_mapel')
 changes[f'staff/{staff}/appRoles']=sorted(set(roles));changes[f'staff/{staff}/unitIds']=sorted(set(d['staff'][staff].get('unitIds',[])+['UNIT-SMK']))
# User requested timetable review, so preserve current homeroom / mentoring assignments.
report={'academicYearId':year,'scheduleCount':count,'homerooms':0,'activeMentoring':0,'issues':warnings,'sourceSha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'sourceSheet':s.title}
pack={'id':update_id,'changes':changes,'report':report}
(ROOT/'seed/imports/smk-2026-2027-update.json').write_text(json.dumps(pack,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False))

for path,value in changes.items():
 parts=path.split('/');target=d
 for part in parts[:-1]:target=target.setdefault(part,{})
 target[parts[-1]]=value
(ROOT/'seed/master-data.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
