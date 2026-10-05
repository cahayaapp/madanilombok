"""Read supplied workbooks without modifying them; preserve cell provenance."""
import json,hashlib
from pathlib import Path
import openpyxl
root=Path(__file__).resolve().parents[1]
result={'files':[],'collections':[]}
for filename,kind,header,teacher,start,cols in [('KELOMPOK HALAOH.xlsx','quran',5,6,7,range(3,11)),('BELAJAR BAHASA ARAB.xlsx','arabic',4,5,6,range(3,6)),('data santri perkamar.xlsx','rooms',5,None,6,range(3,9))]:
 p=Path('/Users/haimac/Downloads')/filename;sheet=openpyxl.load_workbook(p,data_only=True).worksheets[0]
 result['files'].append({'name':filename,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'sheet':sheet.title})
 for col in cols:
  rows=[{'name':str(sheet.cell(r,col).value).strip(),'cell':sheet.cell(r,col).coordinate} for r in range(start,sheet.max_row+1) if sheet.cell(r,col).value is not None]
  result['collections'].append({'key':f'{kind}-P-{col}','kind':kind,'gender':'P','name':str(sheet.cell(header,col).value).strip(),'mentor':str(sheet.cell(teacher,col).value or '').strip() if teacher else '', 'source':filename,'members':rows,'staffRoom':kind=='rooms' and col==4})
(root/'seed/imports/boarding-2026-workbooks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
